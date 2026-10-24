"""Resume upload validation, text extraction (PDF/DOCX) and heuristic structuring.

The parsed structure is a *proposal* shown to the user for review; it only
becomes part of the verified profile after the user saves/confirms it.
"""
from __future__ import annotations

import io
import re
import zipfile
from typing import Optional

from .textutil import find_skills

PDF_MAGIC = b"%PDF"
ZIP_MAGIC = b"PK\x03\x04"

SECTION_ALIASES = {
    "summary": ["summary", "professional summary", "profile", "about", "objective", "about me"],
    "experience": ["experience", "work experience", "professional experience", "employment",
                   "employment history", "work history", "career history"],
    "education": ["education", "academic background", "education & training"],
    "skills": ["skills", "technical skills", "core skills", "core competencies", "competencies",
               "technologies", "tools", "skills & tools", "technical proficiencies", "key skills"],
    "projects": ["projects", "selected projects", "personal projects", "key projects", "academic projects"],
    "publications": ["publications", "research publication", "research publications", "publication"],
    "achievements": ["achievements", "awards", "honors", "awards & achievements", "leadership", "activities"],
    "certifications": ["certifications", "certificates", "licenses", "licenses & certifications",
                       "certifications & licenses"],
}

DATE_TOKEN = r"(?:(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+\d{4}|\d{1,2}/\d{4}|\d{4})"
DATE_RANGE_RE = re.compile(
    rf"(?P<start>{DATE_TOKEN})\s*(?:-|–|—|to)\s*(?P<end>{DATE_TOKEN}|present|current|now)", re.I
)
BULLET_RE = re.compile(r"^\s*(?:[•●▪◦‣∙·\-\*–]|\d+[.)])\s+")
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(\+?\d[\d\s().-]{7,}\d)")
URL_RE = re.compile(r"(https?://\S+|(?:www\.)?(?:linkedin\.com|github\.com)/\S+)", re.I)


COMPANY_HINT = re.compile(r"\b(inc|llc|ltd|limited|corp|corporation|pvt|plc|gmbh|co\.|company|solutions|technologies|"
                          r"labs|group|university|college|institute|bank|consulting)\b", re.I)
DATE_ONLY_RE = re.compile(r"^(?:(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+)?\d{4}$|^(?:present|current)$", re.I)


class UploadError(ValueError):
    pass


def validate_upload(filename: str, data: bytes, max_bytes: int) -> str:
    """Return a normalized kind ('pdf'|'docx') or raise UploadError."""
    if not data:
        raise UploadError("The file is empty.")
    if len(data) > max_bytes:
        raise UploadError(f"File is too large (max {max_bytes // (1024 * 1024)} MB).")
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        if not data.startswith(PDF_MAGIC):
            raise UploadError("File has a .pdf extension but is not a valid PDF.")
        return "pdf"
    if name.endswith(".docx"):
        if not data.startswith(ZIP_MAGIC):
            raise UploadError("File has a .docx extension but is not a valid DOCX.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = z.namelist()
                if "word/document.xml" not in names:
                    raise UploadError("DOCX is missing word/document.xml.")
                # zip-bomb guard
                if sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024:
                    raise UploadError("DOCX expands to an unreasonable size.")
        except zipfile.BadZipFile as e:
            raise UploadError("Corrupt DOCX file.") from e
        return "docx"
    raise UploadError("Only PDF and DOCX resumes are supported.")


def extract_text(kind: str, data: bytes) -> str:
    if kind == "pdf":
        import pymupdf  # PyMuPDF

        with pymupdf.open(stream=data, filetype="pdf") as doc:
            if doc.needs_pass:
                raise UploadError("Password-protected PDFs are not supported.")
            text = "\n".join(page.get_text("text") for page in doc)
    elif kind == "docx":
        import docx  # python-docx

        d = docx.Document(io.BytesIO(data))
        lines = []
        for p in d.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if "list" in style and not BULLET_RE.match(t):
                t = "• " + t
            lines.append(t)
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(dict.fromkeys(cells)))
        text = "\n".join(lines)
    else:  # pragma: no cover
        raise UploadError("Unsupported file type")
    text = text.replace(" ", " ").replace("\r", "")
    text = re.sub(r"[ \t]+", " ", text)
    if len(text.strip()) < 30:
        raise UploadError(
            "Could not extract text. If this is a scanned PDF, export a text-based PDF or DOCX instead."
        )
    return text.strip()


def _heading_key(line: str) -> Optional[str]:
    clean = re.sub(r"[^a-z& ]", "", line.lower()).strip()
    if not clean or len(clean) > 40:
        return None
    for key, aliases in SECTION_ALIASES.items():
        if clean in aliases:
            return key
    return None


def split_sections(text: str) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        key = _heading_key(line)
        if key:
            current = key
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return sections


def _parse_experience(lines: list[str]) -> list[dict]:
    """Group lines into roles using date ranges as anchors."""
    roles: list[dict] = []
    pending_header: list[str] = []
    cur: Optional[dict] = None
    for line in lines:
        m = DATE_RANGE_RE.search(line)
        is_bullet = bool(BULLET_RE.match(line))
        if m and not is_bullet:
            head_text = (line[: m.start()] + line[m.end():]).strip(" |,–—-\t")
            header_bits = [h for h in pending_header if h]
            # non-bullet lines after the previous role's bullets are this role's header
            if cur is not None and cur.get("_trailing"):
                header_bits = cur.pop("_trailing") + header_bits
            if head_text:
                header_bits.append(head_text)
            title, company, location = "", "", ""
            parts: list[str] = []
            for h in header_bits:
                parts.extend([p.strip() for p in re.split(r"\s+[|@·•—–]\s+|\s+at\s+|,\s+(?=[A-Z])", h) if p.strip()])
            # "Company — Title" vs "Title — Company": put the company-looking part second
            if len(parts) >= 2 and COMPANY_HINT.search(parts[0]) and not COMPANY_HINT.search(parts[1]):
                parts[0], parts[1] = parts[1], parts[0]
            if parts:
                title = parts[0]
            if len(parts) > 1:
                company = parts[1]
            if len(parts) > 2:
                location = ", ".join(parts[2:])
            cur = {"title": title, "company": company, "location": location,
                   "start": m.group("start").title(), "end": m.group("end").title(), "bullets": []}
            roles.append(cur)
            pending_header = []
            continue
        if is_bullet:
            if cur is None:
                continue
            if cur.get("_trailing"):
                # trailing text was a wrapped bullet continuation
                if cur["bullets"]:
                    cur["bullets"][-1] += " " + " ".join(cur.pop("_trailing"))
                else:
                    cur.pop("_trailing")
            cur["bullets"].append(BULLET_RE.sub("", line).strip())
            continue
        # plain line: either continuation of a bullet, or a header for the next role
        if cur is None:
            pending_header.append(line)
        else:
            cur.setdefault("_trailing", []).append(line)
            # keep at most two candidate header lines
            cur["_trailing"] = cur["_trailing"][-2:] if len(cur["_trailing"]) > 2 else cur["_trailing"]
    for r in roles:
        trailing = r.pop("_trailing", None)
        if trailing and r["bullets"]:
            r["bullets"][-1] += " " + " ".join(trailing)
    return roles


def _parse_education(lines: list[str]) -> list[dict]:
    out: list[dict] = []
    cur: Optional[dict] = None
    degree_re = re.compile(r"\b(b\.?s\.?|b\.?a\.?|m\.?s\.?|m\.?a\.?|mba|ph\.?d|bachelor|master|associate|diploma|doctor|b\.?tech|m\.?tech|b\.?e\.?)\b", re.I)
    for line in lines:
        text = BULLET_RE.sub("", line).strip()
        m = DATE_RANGE_RE.search(text) or re.search(r"\b(19|20)\d{2}\b", text)
        if degree_re.search(text):
            if cur and not cur["degree"]:
                cur["degree"] = text
            else:
                cur = {"institution": "", "degree": text, "field": "", "start": "", "end": "", "details": ""}
                out.append(cur)
        elif re.search(r"(university|college|institute|school|academy)", text, re.I):
            if cur and not cur["institution"]:
                cur["institution"] = text
            else:
                cur = {"institution": text, "degree": "", "field": "", "start": "", "end": "", "details": ""}
                out.append(cur)
        elif cur:
            cur["details"] = (cur["details"] + " " + text).strip()
        if cur and m and not cur["end"]:
            if hasattr(m, "groupdict") and m.groupdict().get("end"):
                cur["start"], cur["end"] = m.group("start"), m.group("end")
            else:
                cur["end"] = m.group(0)
    return out


def _split_list(lines: list[str]) -> list[str]:
    items: list[str] = []
    for line in lines:
        line = BULLET_RE.sub("", line)
        if ":" in line and len(line.split(":", 1)[0]) < 30:
            line = line.split(":", 1)[1]
        for part in re.split(r"[,;|•·]", line):
            p = part.strip(" .")
            if 1 < len(p) <= 60:
                items.append(p)
    return list(dict.fromkeys(items))


def parse_resume_text(text: str) -> dict:
    sections = split_sections(text)
    header = sections.get("header", [])
    email = next((m.group(0) for m in map(EMAIL_RE.search, header + text.splitlines()[:15]) if m), "")
    phone = ""
    for line in header[:10]:
        m = PHONE_RE.search(line)
        if m and len(re.sub(r"\D", "", m.group(1))) >= 9:
            phone = m.group(1).strip()
            break
    links = list(dict.fromkeys(m.group(0).rstrip(").,") for line in header[:12] for m in URL_RE.finditer(line)))
    name = ""
    for line in header[:4]:
        if not EMAIL_RE.search(line) and not PHONE_RE.search(line) and 1 <= len(line.split()) <= 5 \
                and not URL_RE.search(line):
            name = line.strip()
            break
    headline = ""
    for line in header[1:5]:
        if line != name and not EMAIL_RE.search(line) and not URL_RE.search(line) and not PHONE_RE.search(line):
            headline = line
            break

    experience = _parse_experience(sections.get("experience", []))
    skills_listed = _split_list(sections.get("skills", []))
    listed_lower = {x.lower() for x in skills_listed}
    detected = []
    for s in find_skills(text):
        if s in listed_lower:
            continue
        m = re.search(r"(?<![A-Za-z0-9])" + re.escape(s) + r"(?![A-Za-z0-9])", text, re.I)
        detected.append(m.group(0) if m else s)  # keep the casing used in the resume
    skills = list(dict.fromkeys(skills_listed + detected))
    projects = []
    in_bullet = False
    for line in sections.get("projects", []):
        t = BULLET_RE.sub("", line).strip()
        if DATE_ONLY_RE.match(t):
            continue  # dates sit on the project's title line in most layouts
        if BULLET_RE.match(line) and projects:
            projects[-1]["description"] = (projects[-1]["description"] + " " + t).strip()
            in_bullet = True
        elif in_bullet and projects and (t[:1].islower() or not projects[-1]["description"].rstrip().endswith((".", "!", "?"))):
            projects[-1]["description"] = (projects[-1]["description"] + " " + t).strip()  # wrapped bullet
        else:
            name_part, _, desc = t.partition(" - ")
            projects.append({"name": name_part.strip()[:200], "description": desc.strip(), "link": ""})
            in_bullet = False

    warnings = []
    if not experience:
        warnings.append("No dated work experience found — please add roles manually.")
    if not email:
        warnings.append("No email address detected.")
    return {
        "full_name": name,
        "email": email,
        "phone": phone,
        "headline": headline,
        "links": links,
        "summary": " ".join(sections.get("summary", []))[:5000],
        "experience": experience,
        "education": _parse_education(sections.get("education", [])),
        "skills": skills[:80],
        "certifications": [BULLET_RE.sub("", c).strip() for c in sections.get("certifications", [])][:40],
        "projects": projects[:30],
        "warnings": warnings,
        "sections_found": [k for k in sections if k != "header"],
    }
