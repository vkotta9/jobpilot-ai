"""ATS-friendly DOCX and PDF rendering for resumes and cover letters.

Single column, standard section headings, real text (no images/tables for
content), common fonts — the format applicant tracking systems parse best.
"""
from __future__ import annotations

import csv
import io
from typing import Iterable
from xml.sax.saxutils import escape

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from .tailoring import PLACEHOLDER_RE


def _contact_line(h: dict) -> str:
    return " | ".join(x for x in [h.get("email"), h.get("phone"), h.get("location"), *(h.get("links") or [])] if x)


def _dates(e: dict) -> str:
    return " – ".join(x for x in [e.get("start"), e.get("end")] if x)


def _clean(s: str) -> str:
    return PLACEHOLDER_RE.sub("", s or "").strip()


# ---------------------------------------------------------------- DOCX


def resume_docx(c: dict) -> bytes:
    d = DocxDocument()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    for s in d.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Pt(50)
    h = c.get("header", {})
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h.get("full_name") or "")
    r.bold, r.font.size = True, Pt(18)
    if c.get("headline"):
        q = d.add_paragraph(c["headline"])
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q = d.add_paragraph(_contact_line(h))
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def heading(t: str) -> None:
        hp = d.add_paragraph()
        hr = hp.add_run(t.upper())
        hr.bold, hr.font.size = True, Pt(11.5)

    if c.get("summary"):
        heading("Summary")
        d.add_paragraph(_clean(c["summary"]))
    if c.get("skills"):
        heading("Skills")
        d.add_paragraph(", ".join(c["skills"]))
    if c.get("experience"):
        heading("Experience")
        for e in c["experience"]:
            ep = d.add_paragraph()
            er = ep.add_run(" — ".join(x for x in [e.get("title"), e.get("company")] if x))
            er.bold = True
            meta = " | ".join(x for x in [e.get("location"), _dates(e)] if x)
            if meta:
                ep.add_run("   " + meta)
            for b in e.get("bullets", []):
                d.add_paragraph(_clean(b["text"] if isinstance(b, dict) else b), style="List Bullet")
    if c.get("projects"):
        heading("Projects")
        for pr in c["projects"]:
            pp = d.add_paragraph()
            pp.add_run(pr.get("name", "")).bold = True
            if pr.get("description"):
                pp.add_run(" — " + pr["description"])
    if c.get("education"):
        heading("Education")
        for ed in c["education"]:
            line = ", ".join(x for x in [ed.get("degree"), ed.get("field"), ed.get("institution")] if x)
            ep = d.add_paragraph(line)
            if ed.get("end"):
                ep.add_run(f"  ({_dates(ed)})")
    if c.get("certifications"):
        heading("Certifications")
        for cert in c["certifications"]:
            d.add_paragraph(cert, style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def letter_docx(c: dict) -> bytes:
    d = DocxDocument()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    for para in _clean(c.get("body", "")).split("\n\n"):
        d.add_paragraph(para)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF

_S = {
    "name": ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=17, leading=21, alignment=1),
    "center": ParagraphStyle("center", fontName="Helvetica", fontSize=9.5, leading=12, alignment=1),
    "h": ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=8, spaceAfter=3),
    "role": ParagraphStyle("role", fontName="Helvetica-Bold", fontSize=10, leading=13, spaceBefore=4),
    "body": ParagraphStyle("body", fontName="Helvetica", fontSize=10, leading=13),
    "letter": ParagraphStyle("letter", fontName="Helvetica", fontSize=11, leading=15, spaceAfter=9),
}


def _pdf(flow: list) -> bytes:
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=0.7 * inch, rightMargin=0.7 * inch,
                      topMargin=0.6 * inch, bottomMargin=0.6 * inch, title="Resume").build(flow)
    return buf.getvalue()


def resume_pdf(c: dict) -> bytes:
    h = c.get("header", {})
    f: list = [Paragraph(escape(h.get("full_name") or ""), _S["name"])]
    if c.get("headline"):
        f.append(Paragraph(escape(c["headline"]), _S["center"]))
    f.append(Paragraph(escape(_contact_line(h)), _S["center"]))

    def sec(title: str) -> None:
        f.append(Paragraph(title.upper(), _S["h"]))

    if c.get("summary"):
        sec("Summary")
        f.append(Paragraph(escape(_clean(c["summary"])), _S["body"]))
    if c.get("skills"):
        sec("Skills")
        f.append(Paragraph(escape(", ".join(c["skills"])), _S["body"]))
    if c.get("experience"):
        sec("Experience")
        for e in c["experience"]:
            title = " — ".join(x for x in [e.get("title"), e.get("company")] if x)
            meta = " | ".join(x for x in [e.get("location"), _dates(e)] if x)
            f.append(Paragraph(escape(title) + (f'<font name="Helvetica" size="9">   {escape(meta)}</font>' if meta else ""), _S["role"]))
            items = [ListItem(Paragraph(escape(_clean(b["text"] if isinstance(b, dict) else b)), _S["body"]), leftIndent=12)
                     for b in e.get("bullets", [])]
            if items:
                f.append(ListFlowable(items, bulletType="bullet", start="•", leftIndent=12))
    if c.get("projects"):
        sec("Projects")
        for pr in c["projects"]:
            f.append(Paragraph(f"<b>{escape(pr.get('name', ''))}</b>" + (f" — {escape(pr['description'])}" if pr.get("description") else ""), _S["body"]))
    if c.get("education"):
        sec("Education")
        for ed in c["education"]:
            line = ", ".join(x for x in [ed.get("degree"), ed.get("field"), ed.get("institution")] if x)
            f.append(Paragraph(escape(line + (f"  ({_dates(ed)})" if ed.get("end") else "")), _S["body"]))
    if c.get("certifications"):
        sec("Certifications")
        for cert in c["certifications"]:
            f.append(Paragraph("• " + escape(cert), _S["body"]))
    f.append(Spacer(1, 4))
    return _pdf(f)


def letter_pdf(c: dict) -> bytes:
    flow = [Paragraph(escape(p).replace("\n", "<br/>"), _S["letter"]) for p in _clean(c.get("body", "")).split("\n\n") if p.strip()]
    return _pdf(flow)


# ---------------------------------------------------------------- CSV


def to_csv(rows: Iterable[dict], columns: list[str]) -> str:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=columns, extrasaction="ignore")
    w.writeheader()
    for r in rows:
        # neutralise spreadsheet formula injection
        w.writerow({k: ("'" + v if isinstance(v, str) and v[:1] in "=+-@" else v) for k, v in r.items()})
    return buf.getvalue()
