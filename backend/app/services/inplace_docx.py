"""Format-preserving resume editing for DOCX files.

Units are paragraphs (summary text, bullet points, skills lines). Editing a
unit rewrites only that paragraph's runs, reusing the formatting of its own
regular and bold runs, so fonts, sizes, colours, list styles, indentation,
tables and everything else in the document stay exactly as the user made it.
"""
from __future__ import annotations

import copy
import io
import re
from dataclasses import dataclass, field

import docx
from docx.oxml.ns import qn

from .inplace_pdf import BULLET_CHARS, EDITABLE_SECTIONS, FitError, _section_key, normspace


@dataclass
class DocxUnit:
    id: str
    kind: str
    section: str
    heading: str
    index: int  # paragraph index in document.paragraphs
    rich: list[tuple[str, bool]] = field(default_factory=list)
    label: str = ""
    chars: int = 0

    @property
    def text(self) -> str:
        return normspace("".join(t for t, _ in self.rich))

    def to_public(self) -> dict:
        return {"id": self.id, "kind": self.kind, "section": self.section, "heading": self.heading,
                "text": self.text, "rich": self.rich, "label": self.label, "lines": None, "page": 0}


def _is_list(p) -> bool:
    ppr = p._p.pPr
    if ppr is not None and ppr.find(qn("w:numPr")) is not None:
        return True
    name = (p.style.name or "").lower() if p.style is not None else ""
    return "list" in name or "bullet" in name


def _run_bold(r) -> bool:
    if r.bold is not None:
        return bool(r.bold)
    try:
        return bool(r.style is not None and r.style.font.bold)
    except Exception:
        return False


def _rich(p) -> list[tuple[str, bool]]:
    out: list[tuple[str, bool]] = []
    for r in p.runs:
        if not r.text:
            continue
        b = _run_bold(r)
        if out and out[-1][1] == b:
            out[-1] = (out[-1][0] + r.text, b)
        else:
            out.append((r.text, b))
    if out and out[0][0].lstrip()[:1] in BULLET_CHARS:
        t = out[0][0].lstrip()[1:].lstrip()
        out[0] = (t, out[0][1])
    return [x for x in out if x[0]]


def _is_heading(p, text: str) -> bool:
    name = (p.style.name or "").lower() if p.style is not None else ""
    if name.startswith("heading") or name == "title":
        return _section_key(text) is not None or name.startswith("heading")
    if len(text) > 45:
        return False
    # a short paragraph that is exactly a known section name is a heading
    return _section_key(text) is not None


def analyze(data: bytes) -> dict:
    d = docx.Document(io.BytesIO(data))
    units: list[DocxUnit] = []
    section, heading, seen = "header", "", False
    for i, p in enumerate(d.paragraphs):
        text = normspace(p.text)
        if not text:
            continue
        if _is_heading(p, text):
            heading, section, seen = text, (_section_key(text) or "other"), True
            continue
        if not seen:
            continue
        bullet = _is_list(p) or text[:1] in BULLET_CHARS
        kind = None
        if section == "skills":
            kind = "skill_line"
        elif bullet and section in EDITABLE_SECTIONS:
            kind = "bullet"
        elif section == "summary":
            kind = "paragraph"
        if not kind:
            continue
        rich = _rich(p)
        if not rich:
            continue
        label = ""
        if kind == "skill_line":
            if rich[0][1] and rich[0][0].rstrip().endswith(":"):
                label = rich[0][0].strip()
            elif ":" in rich[0][0][:40]:
                label = rich[0][0].split(":", 1)[0].strip() + ":"
        units.append(DocxUnit(id=f"u{len(units)}", kind=kind, section=section, heading=heading, index=i, rich=rich,
                              label=label, chars=len(text)))
    return {"units": units}


def _clone_run(src_run, text: str, bold: bool | None = None):
    new = copy.deepcopy(src_run._r)
    for t in new.findall(qn("w:t")):
        new.remove(t)
    for tag in ("w:br", "w:tab", "w:drawing", "w:fldChar", "w:instrText"):
        for el in new.findall(qn(tag)):
            new.remove(el)
    from docx.text.run import Run

    run = Run(new, src_run._parent)
    run.text = text
    if bold is not None:
        run.bold = bold
    return new


def render(data: bytes, units: list[DocxUnit], edits: dict[str, list[tuple[str, bool]]],
           skill_additions: dict[str, list[str]]) -> tuple[bytes, dict[str, list[str]]]:
    d = docx.Document(io.BytesIO(data))
    paras = d.paragraphs
    by_id = {u.id: u for u in units}
    warnings: dict[str, list[str]] = {}
    for uid, runs in edits.items():
        u = by_id[uid]
        p = paras[u.index]
        existing = [r for r in p.runs if r.text]
        if not existing:
            raise FitError(f"{uid}: paragraph has no text runs")
        reg_src = next((r for r in existing if not _run_bold(r) and r.text.strip()), existing[0])
        bold_src = next((r for r in existing if _run_bold(r) and r.text.strip()), None)
        prefix = ""
        first = existing[0].text
        if first.lstrip()[:1] in BULLET_CHARS:  # manual bullet glyph typed in the text
            prefix = first[: len(first) - len(first.lstrip())] + first.lstrip()[0] + " "
        anchor = existing[0]._r
        new_elems = []
        if prefix:
            new_elems.append(_clone_run(existing[0], prefix, None))
        for text, bold in runs:
            if bold and bold_src is not None:
                new_elems.append(_clone_run(bold_src, text, None))
            else:
                new_elems.append(_clone_run(reg_src, text, True if bold else None if not _run_bold(reg_src) else False))
        for el in new_elems:
            anchor.addprevious(el)
        for r in existing:
            r._r.getparent().remove(r._r)
        warnings[uid] = []
    for uid, items in skill_additions.items():
        if not items or uid in edits:
            continue
        u = by_id[uid]
        p = paras[u.index]
        existing = [r for r in p.runs if r.text]
        src = next((r for r in reversed(existing) if not _run_bold(r)), existing[-1])
        existing[-1]._r.addnext(_clone_run(src, ", " + ", ".join(items), None))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue(), warnings


def full_text(data: bytes) -> str:
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)
