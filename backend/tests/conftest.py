import io
import os
import sys
import tempfile
import uuid
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="jobpilot-test-")
os.environ["JOBPILOT_DATA_DIR"] = _TMP
os.environ["DATABASE_URL"] = os.environ.get("TEST_DATABASE_URL") or f"sqlite:///{_TMP}/test.db"
os.environ["JOBPILOT_SCHEDULER"] = "0"
os.environ["JOBPILOT_SOURCE_PAUSE"] = "0"
os.environ["OPENAI_API_KEY"] = ""
os.environ["JOBPILOT_ALLOW_MOCK_SOURCES"] = "1"
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402
from app.security import limiter  # noqa: E402


@pytest.fixture(scope="session")
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture(autouse=True)
def _reset_limits():
    limiter.reset()
    yield


def make_user(client):
    email = f"u{uuid.uuid4().hex[:8]}@example.com"
    r = client.post("/api/auth/register", json={"email": email, "password": "correct-horse-1"})
    assert r.status_code == 200, r.text
    client.cookies.clear()  # use bearer tokens explicitly in tests
    return {"Authorization": f"Bearer {r.json()['token']}"}, email


@pytest.fixture
def user(client):
    return make_user(client)[0]


SAMPLE_PROFILE = {
    "full_name": "Jordan Rivera",
    "email": "jordan@example.com",
    "phone": "+1 555 010 2030",
    "location": "New York, NY",
    "headline": "Backend Engineer",
    "summary": "",
    "links": ["https://github.com/jordan"],
    "skills": ["Python", "FastAPI", "PostgreSQL", "Docker", "AWS", "Redis"],
    "experience": [
        {"company": "Acme Corp", "title": "Senior Software Engineer", "location": "New York, NY", "start": "Jan 2020", "end": "Present",
         "bullets": ["responsible for the payments API in Python and FastAPI serving 2M requests/day",
                     "Cut p95 latency by 40% by adding Redis caching",
                     "Mentored 3 junior engineers"]},
        {"company": "Beta LLC", "title": "Software Engineer", "location": "Boston, MA", "start": "Jun 2016", "end": "Dec 2019",
         "bullets": ["Built ETL jobs on PostgreSQL", "Containerised services with Docker and deployed to AWS"]},
    ],
    "education": [{"institution": "State University", "degree": "B.S. Computer Science", "field": "", "start": "", "end": "2016", "details": ""}],
    "certifications": ["AWS Certified Developer"],
    "projects": [],
    "preferences": {"target_titles": ["Backend Engineer"], "seniority": ["senior"], "workplace": ["remote", "hybrid"],
                    "locations": ["New York"], "salary_min": 140000},
    "work_authorization": {"authorized_countries": ["United States"], "requires_sponsorship": None},
    "confirm": True,
}


def setup_profile(client, headers, **overrides):
    body = {**SAMPLE_PROFILE, **overrides}
    r = client.put("/api/profile", json=body, headers=headers)
    assert r.status_code == 200, r.text
    return r.json()


def sample_docx_bytes() -> bytes:
    import docx

    d = docx.Document()
    for line in ["Jordan Rivera", "Backend Engineer", "jordan@example.com | +1 555 010 2030 | github.com/jordan",
                 "Summary", "Backend engineer who builds reliable APIs.",
                 "Experience", "Senior Software Engineer | Acme Corp | Jan 2020 – Present"]:
        d.add_paragraph(line)
    for b in ["Built the payments API in Python and FastAPI", "Cut p95 latency by 40% with Redis caching"]:
        d.add_paragraph(b, style="List Bullet")
    d.add_paragraph("Software Engineer | Beta LLC | Jun 2016 – Dec 2019")
    d.add_paragraph("Built ETL jobs on PostgreSQL", style="List Bullet")
    for line in ["Education", "B.S. Computer Science", "State University, 2016", "Skills", "Python, FastAPI, PostgreSQL, Docker, AWS"]:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def sample_pdf_bytes() -> bytes:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=LETTER)
    y = 740
    for line in ["Jordan Rivera", "jordan@example.com", "EXPERIENCE", "Senior Software Engineer | Acme Corp | Jan 2020 - Present",
                 "• Built the payments API in Python", "• Cut latency by 40%", "EDUCATION", "B.S. Computer Science, State University 2016",
                 "SKILLS", "Python, Docker, AWS"]:
        c.drawString(72, y, line)
        y -= 18
    c.save()
    return buf.getvalue()
